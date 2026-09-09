"""Download ATB vintages and retain their release-year-minus-two estimates."""

import hashlib
import json
import re
import shutil

import numpy as np
import openpyxl
import pandas as pd

from atb_config import raw_file_path, resolve_atb_path
from historical_atb import PARAMETERS, archive_path, technology_series
from downloads import download_file


WORKBOOK_METRICS = {
    'Overnight Capital Cost': 'capcost',
    'Fixed Operation and Maintenance Expenses': 'fom',
    'Variable Operation and Maintenance Expenses': 'vom',
    'Net Capacity Factor': 'cf_improvement',
    'Heat Rate': 'heatrate',
}
UNITS = {'capcost': 'USD/kW', 'capcost_energy': 'USD/kWh',
         'fom': 'USD/kW-yr', 'fom_energy': 'USD/kWh-yr',
         'vom': 'USD/MWh', 'heatrate': 'MMBtu/MWh',
         'cf_improvement': 'fraction', 'rte': 'fraction'}


def record(target, metric, value, vintage, source, location, detail, notes=''):
    if target[0] in ('wind-ons', 'wind-ofs') and vintage <= 2022:
        notes = (f'{notes} Resource-class history uses the vintage-specific turbine design; '
                 'it does not assume the ATB 2024 turbine dimensions.').strip()
    if 2022 <= vintage <= 2023 and target[1] in ('Gas-CC_H_2x1', 'Gas-CC_H_2x1-CCS_mod'):
        assumption_row = 14 if target[0] == 'gas' else (18 if vintage == 2022 else 16)
        notes = (f'{notes} H-frame maps to 2x1 per ATB {vintage} workbook '
                 f'Natural Gas_FE!K{assumption_row} and '
                 f'https://atb.nlr.gov/electricity/{vintage}/fossil_energy_technologies; '
                 'the separate 1x1 design was introduced in ATB 2024.').strip()
    return dict(technology=target[0], series=target[1], metric=metric,
                year=vintage - 2, atb_year=vintage, dollar_year=source['dollar_year'],
                value=float(value), unit=UNITS[metric], scenario='Moderate',
                source_file=source['filename'], source_url=source['url'],
                source_location=location, source_technology=detail, notes=notes)


def extract_workbook(path, vintage, source):
    """Read cached cells; macros and formulas are never executed."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        if vintage >= 2017:
            title = str(workbook[workbook.sheetnames[0]]['E2'].value)
            if str(vintage) not in title:
                raise ValueError(f'{path.name}: expected ATB {vintage}, found {title!r}.')
        for sheet in workbook:
            if sheet.title not in {
                'Conventional - Coal', 'Conventional - Gas', 'Conventional - Biopower',
                'Conventional - Nuclear', 'Coal', 'Natural Gas', 'Coal_FE',
                'Natural Gas_FE', 'Nuclear', 'Biopower', 'Solar - CSP',
                'Solar - Utility PV', 'Land-Based Wind', 'Offshore Wind',
            }:
                continue
            rows = list(sheet.iter_rows(max_col=20, values_only=True))
            basis = [value for row in rows[:20] for j, cell in enumerate(row[:-4])
                     if str(cell).strip() in ('Basis Year:', 'Base Year:')
                     for value in row[j + 1:j + 5]
                     if isinstance(value, (int, float)) and 1900 <= value <= 2100]
            if basis != [vintage - 2]:
                raise ValueError(f'{path.name}/{sheet.title}: unexpected base year {basis}.')
            metric = None
            base_column = None
            for row_number, row in enumerate(rows, 1):
                if vintage <= 2020 and any('Future Projections' in str(v) for v in row[:10]):
                    break
                if vintage >= 2021:
                    year_columns = [i for i in range(11, 19)
                                    if row[i] == vintage - 2 and row[i + 1] == vintage - 1]
                    if year_columns:
                        base_column = year_columns[0]
                        metric = None
                        continue
                    if base_column is None:
                        continue
                else:
                    base_column = 11
                if row[9] is not None:
                    metric = next((m for prefix, m in WORKBOOK_METRICS.items()
                                   if str(row[9]).startswith(prefix)), None)
                if metric is None or not isinstance(row[10], str):
                    continue
                if vintage == 2021 and not re.search(r'-\s*Moderate$', row[10]):
                    continue
                if vintage >= 2022 and row[11] != 'Moderate':
                    continue
                value = row[base_column]
                if not isinstance(value, (int, float)) or not np.isfinite(value):
                    continue
                if metric in ('capcost', 'heatrate', 'cf_improvement') and value <= 0:
                    continue
                target = technology_series(sheet.title, row[10], vintage)
                if target:
                    cell = f'{sheet.title}!{openpyxl.utils.get_column_letter(base_column + 1)}{row_number}'
                    records.append(record(target, metric, value, vintage, source, cell, row[10]))
    finally:
        workbook.close()
    return records


def extract_flat(path, vintage, source):
    data = pd.read_csv(path, low_memory=False)
    if set(pd.to_numeric(data.atb_year).unique()) != {vintage}:
        raise ValueError(f'{path.name} does not contain ATB {vintage}.')
    cases = data.groupby('technology').core_metric_case.agg(set).map(
        lambda available: next((name for name in ('Market', 'Exp', 'R&D')
                                if name in available), None)
    )
    if cases.isna().any():
        raise ValueError(f'{path.name}: unrecognized ATB financial cases.')
    data = data.loc[
        pd.to_numeric(data.core_metric_variable, errors='coerce').eq(vintage - 2)
        & data.core_metric_case.eq(data.technology.map(cases))
        & pd.to_numeric(data.crpyears, errors='coerce').eq(30)
        & data.scenario.eq('Moderate')
        & data.core_metric_parameter.isin(PARAMETERS)
    ]
    records = []
    for index, row in data.iterrows():
        detail = row.get('display_name', row.techdetail)
        target = technology_series(row.technology, detail, vintage)
        if not target:
            continue
        metric = PARAMETERS[row.core_metric_parameter]
        value = float(row.value)
        unit = str(row.units).lower()
        expected = {'capcost': '$/kw', 'fom': '$/kw-yr', 'vom': '$/mwh',
                    'heatrate': 'mmbtu/mwh', 'cf_improvement': '%', 'rte': '%'}[metric]
        if unit not in (expected, 'nan'):
            raise ValueError(f'{path.name}: unexpected {metric} unit {row.units!r}.')
        # ATBe labels CF as percent, but stores fractions.
        if metric in ('cf_improvement', 'rte') and not 0 < value <= 1:
            raise ValueError(f'{path.name}: invalid {metric} fraction {value}.')
        if not np.isfinite(value) or value < 0:
            raise ValueError(f'{path.name}: invalid {metric} value {value}.')
        records.append(record(target, metric, value, vintage, source,
                              f'CSV row {index + 2}', detail,
                              f'Case: {row.core_metric_case}. ' + (
                                  'Unit from ATB parameter definition; CSV unit blank.' if unit == 'nan' else '')))
    return records


def extract_battery(path, vintage, source):
    from battery_workbook import extract_battery_costs
    base_year = vintage - 2
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next(workbook[name] for name in (
            'Utility-Scale Battery Storage', 'Utility-Scale Battery - Expand', 'Storage'
        ) if name in workbook.sheetnames)
        sheet_name = sheet.title
        rows = list(sheet.iter_rows(max_col=20, values_only=True))
    finally:
        workbook.close()
    header = next((row for row in rows if any(
        row[i] == base_year and row[i + 1] == base_year + 1
        for i in range(len(row) - 1))), None)
    if header is None:
        raise ValueError(f'{path.name}: battery base year {base_year} is absent.')
    base_column = header.index(base_year)
    if vintage == 2020:
        costs = {row[10]: (n, row[base_column]) for n, row in enumerate(rows, 1)
                 if row[10] in ('2Hr Battery Storage - Moderate', '4Hr Battery Storage - Moderate')
                 and n < 17}
        two_row, two = costs['2Hr Battery Storage - Moderate']
        four_row, four = costs['4Hr Battery Storage - Moderate']
        energy = (four - two) / 2
        values = {'capcost': two - 2 * energy, 'capcost_energy': energy}
        location = f'{sheet_name}!L{two_row},L{four_row}'
        notes = ('Calculated battery components: E=(C4-C2)/(4-2); P=C2-2*E; '
                 'C(h)=P+h*E from ATB 2-hour and 4-hour system costs.')
    else:
        components = extract_battery_costs(path)
        if base_year not in components:
            raise ValueError(f'{path.name}: battery base year {base_year} is absent.')
        values = components.loc[components.Scenario.eq('Moderate')].set_index('cost')[base_year].to_dict()
        notes = ''
    records = []
    for metric, value in values.items():
        if vintage != 2020:
            location = f'{sheet_name}: {metric}, Moderate, {base_year}'
        if not np.isfinite(value) or value <= 0:
            raise ValueError(f'{path.name}: invalid battery {metric} for {base_year}.')
        records.append(record(('battery', 'battery_li'), metric, value,
                              vintage, source, location, 'Battery components', notes))
        records.append(record(('battery', 'battery_li'), metric.replace('capcost', 'fom'),
                              value * 0.025, vintage, source, location,
                              'Battery components', (notes + ' ReEDS FOM convention: 2.5% of capital component.').strip()))
    metric = None
    label_column = 10 if vintage == 2020 else 4
    for n, row in enumerate(rows, 1):
        year_columns = [i for i in range(len(row) - 1)
                        if row[i] == base_year and row[i + 1] == base_year + 1]
        if year_columns:
            base_column = year_columns[0]
            metric = None
            continue
        label = row[label_column - 1]
        if label is not None:
            metric = {'Variable Operation and Maintenance Expenses ($/MWh)': 'vom',
                      'Round-Trip Efficiency': 'rte'}.get(str(label).strip())
        detail = str(row[label_column]).strip()
        if metric is None or '2Hr' not in detail:
            continue
        if not (detail.endswith(' - Moderate') if vintage <= 2021 else row[label_column + 1] == 'Moderate'):
            continue
        value = row[base_column]
        if not isinstance(value, (int, float)) or not np.isfinite(value) or value < 0:
            raise ValueError(f'{path.name}: invalid battery {metric} for {base_year}.')
        if metric == 'rte' and not 0 < value <= 1:
            raise ValueError(f'{path.name}: invalid battery efficiency {value}.')
        location = f'{sheet_name}!{openpyxl.utils.get_column_letter(base_column + 1)}{n}'
        records.append(record(('battery', 'battery_li'), metric, value,
                              vintage, source, location, detail))
    if {row['metric'] for row in records} != {'capcost', 'capcost_energy', 'fom', 'fom_energy', 'vom', 'rte'}:
        raise ValueError(f'{path.name}: incomplete battery base-year parameters.')
    return records


def scrape(config, force=False, no_download=False):
    options = config['historical_atb']
    directory = resolve_atb_path(options['directory'])
    directory.mkdir(parents=True, exist_ok=True)
    releases = {int(year): entries for year, entries in options['releases'].items()
                if int(year) < config['atb']['year']}
    manifest_path = directory / 'source_manifest.csv'
    if manifest_path.is_file():
        previous = pd.read_csv(manifest_path)
        for vintage, entries in previous.groupby('atb_year'):
            if int(vintage) < config['atb']['year'] and int(vintage) not in releases:
                releases[int(vintage)] = entries[
                    ['filename', 'url', 'dollar_year', 'format']
                ].to_dict('records')
    prepared_config = config['historical_data']
    prepared_path = resolve_atb_path(prepared_config['directory']) / prepared_config['atb_filename']
    if prepared_path.is_file():
        prepared = pd.read_csv(prepared_path, keep_default_na=False)
        for vintage, entries in prepared.groupby('atb_year'):
            if int(vintage) >= config['atb']['year'] or int(vintage) in releases:
                continue
            sources = {}
            for encoded in entries.sources:
                source = json.loads(encoded)[0]
                sources[source['file']] = dict(filename=source['file'], url=source['url'],
                                               dollar_year=source['dollar_year'], format=source['format'])
            releases[int(vintage)] = list(sources.values())
    # Retain the current release for the next annual preparation run.
    current = config['atb']['year']
    releases[current] = []
    for kind in ('flat_file', 'workbook'):
        original = raw_file_path(config, kind)
        if not original.is_file():
            raise FileNotFoundError(f'Current ATB input missing: {original}. Run future_atb_scraper.py first.')
        filename = f'atb_{current}_{kind}{original.suffix}'
        destination = directory / filename
        if original.resolve() != destination.resolve():
            shutil.copy2(original, destination)
        releases[current].append(dict(
            filename=filename, url=config['raw_data'][kind]['url'],
            dollar_year=config['atb']['dollar_year'],
            format='flat' if kind == 'flat_file' else 'battery',
        ))
    records, manifest = [], []
    for vintage, entries in sorted(releases.items()):
        for source in entries:
            path = directory / source['filename']
            if not no_download:
                download_file(source['url'], path, force=force and vintage != current,
                              allow_insecure_ssl_fallback=config['raw_data'].get('allow_insecure_ssl_fallback', False))
            if not path.is_file():
                raise FileNotFoundError(f'Archived ATB input missing: {path}')
            extract = {'workbook': extract_workbook, 'flat': extract_flat,
                       'battery': extract_battery}[source['format']]
            extracted = extract(path, vintage, source)
            if not extracted:
                raise ValueError(f'No base-year estimates extracted from {path}.')
            records.extend(extracted)
            if source['format'] == 'workbook' and vintage >= 2020:
                records.extend(extract_battery(path, vintage, source))
            manifest.append(dict(atb_year=vintage, base_year=vintage - 2, **source,
                                 sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    data = pd.DataFrame(records)
    keys = ['technology', 'series', 'metric', 'year']
    if data.duplicated(keys).any():
        raise ValueError(f'Duplicate archived estimates:\n{data.loc[data.duplicated(keys, keep=False), keys]}')
    data = data.sort_values(keys).reset_index(drop=True)
    temporary = archive_path(config).with_suffix('.csv.part')
    data.to_csv(temporary, index=False, lineterminator='\n')
    temporary.replace(archive_path(config))
    pd.DataFrame(manifest).to_csv(directory / 'source_manifest.csv', index=False, lineterminator='\n')
    print(f'Saved {len(data)} ATB base-year estimates to {archive_path(config)}')
    return data

